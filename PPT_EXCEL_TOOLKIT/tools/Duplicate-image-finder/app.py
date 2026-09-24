import collections
import io
import base64
import gc
import cv2
import numpy as np
import streamlit as st
from pptx import Presentation
from PIL import Image

st.set_page_config(page_title="Enterprise Image Matching System", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
    <style>
    [data-testid="stToolbar"] {display: none !important;}
    footer {visibility: hidden !important;}
    [data-testid="stSidebarCollapseButton"] {display: none !important;}
    button[title="Collapse sidebar"] {display: none !important;}
    
    /* Strict Spacing Adjustments */
    .block-container {padding-top: 1rem; padding-bottom: 0rem;}
    [data-testid="stSidebar"] {padding-top: 0.5rem;}
    [data-testid="stSidebar"] .element-container {margin-bottom: 0.2rem !important;}
    [data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p {margin-bottom: 0.1rem !important;}
    
    h1 {font-size: 1.6rem !important; margin-bottom: 0.5rem;}
    h4 {font-size: 1.05rem !important; margin-top: 0.3rem !important; margin-bottom: 0.2rem !important;}
    .stAlert {padding: 0.4rem 0.8rem; margin-bottom: 0.3rem;}
    hr {margin-top: 0.4rem !important; margin-bottom: 0.4rem !important;}
    
    .zoom-img {
        width: 90px;
        border-radius: 3px;
        cursor: pointer;
        transition: transform 0.2s ease;
    }
    .zoom-img:focus {
        position: fixed;
        top: 50%;
        left: 50%;
        transform: translate(-50%, -50%) scale(4);
        max-width: 80vw;
        max-height: 80vh;
        z-index: 99999;
        box-shadow: 0 0 15px rgba(0,0,0,0.7);
        outline: none;
    }

    /* Score Highlighting Badge (>=60) */
    .score-badge-high {
        background-color: #1e4620;
        color: #4caf50;
        padding: 2px 8px;
        border-radius: 4px;
        border: 1px solid #4caf50;
        font-weight: bold;
    }
    .score-badge-normal {
        color: #b0bec5;
    }
    </style>
""", unsafe_allow_html=True)

st.title("Enterprise Image Matching System")

def get_thumbnail_base64(img, max_size=(300, 300)):
    thumb = img.copy()
    if thumb.mode in ("RGBA", "P"):
        thumb = thumb.convert("RGB")
    thumb.thumbnail(max_size)
    buffered = io.BytesIO()
    thumb.save(buffered, format="JPEG", quality=75)
    img_str = base64.b64encode(buffered.getvalue()).decode()
    return f"data:image/jpeg;base64,{img_str}"

def pil_to_cv2(pil_img):
    if pil_img.mode in ("RGBA", "P"):
        pil_img = pil_img.convert("RGB")
    open_cv_image = np.array(pil_img)
    return cv2.cvtColor(open_cv_image, cv2.COLOR_RGB2GRAY)

def get_features(cv_img):
    orb = cv2.ORB_create(nfeatures=500)
    kp, des = orb.detectAndCompute(cv_img, None)
    return des

def compare_features(des1, des2):
    if des1 is None or des2 is None:
        return 0
    bf = cv2.BFMatcher(cv2.NORM_HAMMING, crossCheck=True)
    matches = bf.match(des1, des2)
    good_matches = [m for m in matches if m.distance < 40]
    return len(good_matches)

def format_score_html(score):
    if score >= 60:
        return f'<span class="score-badge-high">Score: {score} features</span>'
    return f'<span class="score-badge-normal">Score: <b>{score}</b> features</span>'

# PPT Extraction with LIVE Progress Tracking
def process_ppt_file_with_progress(ppt_file, label_prefix="Main PPT"):
    prs = Presentation(ppt_file)
    ppt_images = []
    total_slides = len(prs.slides)
    
    progress_bar = st.progress(0)
    status_text = st.empty()
    
    for slide_index, slide in enumerate(prs.slides):
        current_slide_num = slide_index + 1
        
        # Live status update on UI
        status_text.text(f"Processing {label_prefix}: Slide {current_slide_num} of {total_slides}...")
        progress_bar.progress(current_slide_num / total_slides)
        
        picture_shapes = []
        for shape in slide.shapes:
            if shape.shape_type == 13: # Shape Picture
                picture_shapes.append((shape.top, shape.left, shape))
        
        picture_shapes.sort(key=lambda x: (x[0], x[1]))
        
        for img_count, (_, _, shape) in enumerate(picture_shapes, start=1):
            try:
                image_bytes = shape.image.blob
                with Image.open(io.BytesIO(image_bytes)) as img:
                    cv_img = pil_to_cv2(img)
                    des = get_features(cv_img)
                    b64_str = get_thumbnail_base64(img)
                    
                    ppt_images.append({
                        "slide": current_slide_num,
                        "img_num": img_count,
                        "descriptors": des,
                        "b64_img": b64_str
                    })
            except Exception:
                continue
                
    # Clear progress elements once done
    progress_bar.empty()
    status_text.empty()
    
    return ppt_images, total_slides

# --- SIDEBAR CONTROLS ---
with st.sidebar:
    st.header("Configuration")
    
    search_mode = st.radio(
        "Select Search Mode:",
        ["Single Image Search", "PPT to PPT Batch Search"],
        index=0
    )
    
    main_ppt_file = st.file_uploader("1. Main PPT File", type=["pptx"])

    if search_mode == "Single Image Search":
        single_img_file = st.file_uploader("2. Reference Image File", type=["jpg", "jpeg", "png", "webp"])
        target_ppt_file = None
    else:
        target_ppt_file = st.file_uploader("2. Target PPT File", type=["pptx"])
        single_img_file = None

    search_clicked = st.button("Start Matching Process", use_container_width=True, type="primary")

# --- MAIN WORKFLOW ---
if main_ppt_file is None:
    st.session_state.clear()
    st.info("Please upload the Main PPT File in the sidebar to proceed.")
else:
    if "main_ppt_images" not in st.session_state or st.session_state.get("main_ppt_name") != main_ppt_file.name:
        main_images, main_slide_count = process_ppt_file_with_progress(main_ppt_file, label_prefix="Main PPT File")
        st.session_state["main_ppt_images"] = main_images
        st.session_state["main_ppt_slides"] = main_slide_count
        st.session_state["main_ppt_name"] = main_ppt_file.name
        st.success(f"Main PPT File processed successfully ({main_slide_count} slides, {len(main_images)} images indexed).")

    main_images = st.session_state["main_ppt_images"]
    main_slide_count = st.session_state["main_ppt_slides"]

    # --- MODE 1: SINGLE IMAGE SEARCH ---
    if search_mode == "Single Image Search":
        st.subheader("Single Image Match Results")

        if search_clicked:
            if single_img_file is None:
                st.warning("Please upload a reference image file in the sidebar.")
            else:
                with Image.open(single_img_file) as target_img:
                    target_cv = pil_to_cv2(target_img)
                    target_des = get_features(target_cv)
                    target_b64 = get_thumbnail_base64(target_img)

                raw_matches = []
                for main_item in main_images:
                    score = compare_features(target_des, main_item["descriptors"])
                    if score >= 35:
                        raw_matches.append((main_item, score))
                
                raw_matches.sort(key=lambda x: x[1], reverse=True)

                if raw_matches:
                    highest_score = raw_matches[0][1]
                    strict_matches = [
                        (item, score) for item, score in raw_matches 
                        if score >= max(40, highest_score * 0.85)
                    ]

                    st.success(f"Matching completed: {len(strict_matches)} exact match(es) identified across {main_slide_count} slides in Main PPT.")
                    c1, c2 = st.columns([1, 6])
                    with c1:
                        st.markdown(f'<img src="{target_b64}" class="zoom-img" tabindex="0">', unsafe_allow_html=True)

                    with c2:
                        for match_item, score in strict_matches:
                            score_html = format_score_html(score)
                            st.markdown(f"• **Slide {match_item['slide']}** → Image #{match_item['img_num']} ({score_html})", unsafe_allow_html=True)
                else:
                    st.error(f"No exact match identified across all {main_slide_count} slides in the Main PPT File.")
        else:
            st.info("Upload the Reference Image and click 'Start Matching Process'.")

    # --- MODE 2: PPT TO PPT SEARCH ---
    else:
        st.subheader("PPT to PPT Batch Matching Report")

        if search_clicked:
            if target_ppt_file is None:
                st.warning("Please upload the Target PPT File in the sidebar.")
            else:
                target_images, target_slide_count = process_ppt_file_with_progress(target_ppt_file, label_prefix="Target PPT File")
                
                if not target_images:
                    st.error("No images found in the Target PPT File.")
                else:
                    found_count = 0
                    for target_item in target_images:
                        raw_matches = []
                        for main_item in main_images:
                            score = compare_features(target_item["descriptors"], main_item["descriptors"])
                            if score >= 35:
                                raw_matches.append((main_item, score))
                        
                        raw_matches.sort(key=lambda x: x[1], reverse=True)
                        
                        if raw_matches:
                            highest_score = raw_matches[0][1]
                            strict_matches = [
                                (item, score) for item, score in raw_matches 
                                if score >= max(40, highest_score * 0.85)
                            ]
                        else:
                            strict_matches = []

                        # Hide non-matching items completely
                        if strict_matches:
                            found_count += 1
                            st.markdown(f"#### Target PPT → Slide {target_item['slide']} (Image #{target_item['img_num']})")
                            c1, c2 = st.columns([1, 6])
                            
                            with c1:
                                st.markdown(f'<img src="{target_item["b64_img"]}" class="zoom-img" tabindex="0">', unsafe_allow_html=True)
                            
                            with c2:
                                for main_match, score in strict_matches:
                                    score_html = format_score_html(score)
                                    st.markdown(f"Matched in **Main PPT** → **Slide {main_match['slide']}** (Image #{main_match['img_num']}) — {score_html}", unsafe_allow_html=True)
                            
                            st.divider()

                    if found_count > 0:
                        st.success(f"Batch execution completed: Successfully matched **{found_count} of {len(target_images)}** images from **{target_slide_count} Target PPT slides** against **{main_slide_count} Main PPT slides**.")
                    else:
                        st.warning("No matching images were found between Target PPT and Main PPT.")
        else:
            st.info("Upload the Target PPT File and click 'Start Matching Process'.")

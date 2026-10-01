import re
import tempfile
from pathlib import Path
import pandas as pd
import streamlit as st
from pptx import Presentation
from .common import read_excel, save_uploaded, ppt_text, dataframe_download, find_column, normalize


def run_tool(max_upload_mb=2048):
    st.subheader("1. Upload files")
    excel = st.file_uploader("Excel Master File", type=["xlsx", "xls", "csv"], key="sm_excel")
    ppt = st.file_uploader("PowerPoint File", type=["pptx"], key="sm_ppt")

    if not excel or not ppt:
        st.info("Upload Excel + PPT to continue.")
        return

    df = read_excel(excel)
    ppt_path = save_uploaded(ppt)
    prs = Presentation(str(ppt_path))

    match_col = find_column(df, ["match", "name", "title", "product", "item", "client", "brand", "sap code"])
    if not match_col:
        st.warning("No obvious matching column was found. The first Excel column will be used.")
        match_col = df.columns[0]

    st.success(f"Excel loaded: {len(df)} rows | PPT loaded: {len(prs.slides)} slides")
    if st.button("🚀 Start Matching", type="primary", key="sm_start"):
        rows = []
        slide_texts = [normalize(ppt_text(s)) for s in prs.slides]

        for idx, value in enumerate(df[match_col].fillna("")):
            needle = normalize(value)
            best_slide = ""
            status = "Not Found"
            reason = ""
            confidence = 0

            if needle:
                exact = [i for i, text in enumerate(slide_texts) if needle and needle in text]
                if exact:
                    best_slide = exact[0] + 1
                    status = "Matched"
                    confidence = 100
                    reason = "Exact normalized text match"
                else:
                    tokens = set(needle.split())
                    best_i, best_score = None, 0
                    for i, text in enumerate(slide_texts):
                        score = sum(1 for t in tokens if t in text)
                        if score > best_score:
                            best_i, best_score = i, score
                    if best_i is not None and tokens and best_score / len(tokens) >= 0.5:
                        best_slide = best_i + 1
                        status = "Review"
                        confidence = round(100 * best_score / len(tokens))
                        reason = "Partial token match"

            rows.append({
                "Original Data": value,
                "Match Status": status,
                "Matched PPT Slide": best_slide,
                "Confidence %": confidence,
                "Match Reason": reason,
                "Review Required": "Yes" if status == "Review" else "No",
            })

        report = pd.DataFrame(rows)
        st.dataframe(report, use_container_width=True)
        st.download_button(
            "⬇️ Download Matching Report",
            dataframe_download(report),
            "smart_matcher_report.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

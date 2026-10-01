import io
import zipfile
from pathlib import Path
import streamlit as st
from pptx import Presentation
from .common import save_uploaded, safe_name


def run_tool(max_upload_mb=2048):
    st.subheader("1. Upload PowerPoint")
    ppt = st.file_uploader("PowerPoint File", type=["pptx"], key="img_ppt")
    if not ppt:
        st.info("Upload PPT to continue.")
        return

    if st.button("🖼️ Extract PPT Images", type="primary", key="img_start"):
        prs = Presentation(str(save_uploaded(ppt)))
        memory = io.BytesIO()
        count = 0

        with zipfile.ZipFile(memory, "w", zipfile.ZIP_DEFLATED) as z:
            for slide_no, slide in enumerate(prs.slides, start=1):
                image_index = 0
                for shape in slide.shapes:
                    if getattr(shape, "shape_type", None) == 13 and getattr(shape, "image", None):
                        image_index += 1
                        count += 1
                        ext = shape.image.ext
                        filename = f"slide_{slide_no:04d}_image_{image_index:02d}.{ext}"
                        z.writestr(filename, shape.image.blob)

        memory.seek(0)
        st.success(f"{count} images extracted.")
        st.download_button(
            "⬇️ Download Images ZIP",
            memory.getvalue(),
            "ppt_images.zip",
            "application/zip"
        )

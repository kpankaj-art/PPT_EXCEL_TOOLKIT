import re
import tempfile
from pathlib import Path
import pandas as pd
import streamlit as st
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from .common import read_excel, save_uploaded, dataframe_download, find_column


WIDTH_ALIASES = [
    "w", "width", "widht", "wd", "board width",
    "width inch", "width inches", "width (inch)", "width (inches)"
]
HEIGHT_ALIASES = [
    "h", "height", "hight", "heigt", "ht", "board height",
    "height inch", "height inches", "height (inch)", "height (inches)"
]
SIZE_ALIASES = ["size", "dimension", "dimensions", "w x h", "width x height"]


def _clean_number(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    try:
        number = float(text)
        if number.is_integer():
            return str(int(number))
        return str(number).rstrip("0").rstrip(".")
    except Exception:
        return text


def _shape_text(shape):
    return getattr(shape, "text", "") or ""


def _is_width_text(text):
    t = text.lower().strip()
    return bool(re.search(r"^\s*(w|width)\s*[:=-]?\s*$", t))


def _is_height_text(text):
    t = text.lower().strip()
    return bool(re.search(r"^\s*(h|height)\s*[:=-]?\s*$", t))


def _set_shape_text(shape, text):
    if not getattr(shape, "has_text_frame", False):
        return False
    shape.text_frame.clear()
    p = shape.text_frame.paragraphs[0]
    p.text = str(text)
    p.alignment = PP_ALIGN.CENTER
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return True


def _align_three_shapes(width_shape, x_shape, height_shape):
    shapes = [width_shape, x_shape, height_shape]
    center_y = sum(s.top + s.height / 2 for s in shapes) / len(shapes)
    for s in shapes:
        s.top = int(center_y - s.height / 2)
        if getattr(s, "has_text_frame", False):
            s.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in s.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER

    ordered = sorted(shapes, key=lambda s: s.left)
    gap = int(0.05 * 914400)
    for prev, cur in zip(ordered, ordered[1:]):
        minimum_left = prev.left + prev.width + gap
        if cur.left < minimum_left:
            cur.left = minimum_left


def _find_separate_size(slide):
    width_shape = x_shape = height_shape = None
    for shape in slide.shapes:
        text = _shape_text(shape)
        if not text:
            continue
        if _is_width_text(text):
            width_shape = shape
        elif text.strip().upper() in {"X", "×"}:
            x_shape = shape
        elif _is_height_text(text):
            height_shape = shape

    if width_shape and x_shape and height_shape:
        return width_shape, x_shape, height_shape
    return None


def _find_x_and_height(slide):
    """Find an X separator and the height value even when the W field is blank."""
    x_shape = None
    height_shape = None
    for shape in slide.shapes:
        text = _shape_text(shape).strip()
        if text.upper() in {"X", "×"}:
            x_shape = shape
        elif text and re.fullmatch(r"\d+(?:\.\d+)?", text):
            # Candidate numeric shape to the right of X, usually the H value.
            if x_shape is not None and shape.left >= x_shape.left:
                if height_shape is None or shape.left < height_shape.left:
                    height_shape = shape

    if x_shape is None or height_shape is None:
        return None
    return x_shape, height_shape


def _add_width_shape(slide, x_shape, height_shape, width_text):
    """Create the missing W text box immediately before X, matching H styling/size."""
    from pptx.util import Inches

    # Put W directly before X with a small gap. Match the height box dimensions.
    gap = Inches(0.05)
    new_width = height_shape.width
    new_height = height_shape.height
    left = max(0, x_shape.left - new_width - gap)
    top = height_shape.top

    new_shape = slide.shapes.add_textbox(left, top, new_width, new_height)

    # Copy basic paragraph/font formatting from H where possible.
    if getattr(height_shape, "has_text_frame", False) and height_shape.text_frame.paragraphs:
        src_p = height_shape.text_frame.paragraphs[0]
        dst_p = new_shape.text_frame.paragraphs[0]
        dst_p.alignment = src_p.alignment
        dst_p.level = src_p.level
        if src_p.runs:
            src_run = src_p.runs[0]
            run = dst_p.add_run()
            run.text = str(width_text)
            run.font.name = src_run.font.name
            run.font.size = src_run.font.size
            run.font.bold = src_run.font.bold
            run.font.italic = src_run.font.italic
            run.font.underline = src_run.font.underline
            if src_run.font.color.type is not None:
                try:
                    run.font.color.rgb = src_run.font.color.rgb
                except Exception:
                    pass
        else:
            dst_p.text = str(width_text)
    else:
        new_shape.text = str(width_text)

    new_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return new_shape


def _find_combined_size(slide):
    for shape in slide.shapes:
        text = _shape_text(shape)
        if re.search(r"\b\d+(?:\.\d+)?\s*[xX×]\s*\d+(?:\.\d+)?\b", text):
            return shape
    return None


def run_tool(max_upload_mb=2048):
    st.subheader("1. Upload files")
    excel = st.file_uploader("Excel Master File", type=["xlsx", "xls", "csv"], key="size_excel")
    ppt = st.file_uploader("PowerPoint File", type=["pptx"], key="size_ppt")

    if not excel or not ppt:
        st.info("Upload Excel + PPT to continue.")
        return

    df = read_excel(excel)
    width_col = find_column(df, WIDTH_ALIASES)
    height_col = find_column(df, HEIGHT_ALIASES)

    if not width_col or not height_col:
        st.error("Excel mein W/Width aur H/Height columns nahi mile.")
        st.write("Detected columns:", list(df.columns))
        return

    st.success(f"Excel loaded: {len(df)} rows | W = {width_col} | H = {height_col}")

    if st.button("🚀 Fix PPT Size", type="primary", key="size_start"):
        ppt_path = save_uploaded(ppt)
        prs = Presentation(str(ppt_path))
        report = []

        for slide_no, slide in enumerate(prs.slides, start=1):
            if slide_no > len(df):
                report.append({"Slide": slide_no, "Status": "No Excel Row"})
                continue

            width = _clean_number(df.iloc[slide_no - 1][width_col])
            height = _clean_number(df.iloc[slide_no - 1][height_col])

            if not width or not height:
                report.append({"Slide": slide_no, "Status": "W/H Missing"})
                continue

            separate = _find_separate_size(slide)
            combined = _find_combined_size(slide)
            partial = _find_x_and_height(slide)

            if separate:
                ws, xs, hs = separate
                _set_shape_text(ws, width)
                _set_shape_text(xs, "X")
                _set_shape_text(hs, height)
                _align_three_shapes(ws, xs, hs)
                report.append({
                    "Slide": slide_no,
                    "Status": "Updated",
                    "Width": width,
                    "Height": height,
                    "Action": f"{width} X {height}"
                })
            elif partial:
                # Common template case: W text box is blank/missing, while X and H exist.
                xs, hs = partial
                ws = _add_width_shape(slide, xs, hs, width)
                _set_shape_text(xs, "X")
                _set_shape_text(hs, height)
                _align_three_shapes(ws, xs, hs)
                report.append({
                    "Slide": slide_no,
                    "Status": "Updated (W field created)",
                    "Width": width,
                    "Height": height,
                    "Action": f"{width} X {height}"
                })
            elif combined:
                _set_shape_text(combined, f"{width} X {height}")
                report.append({
                    "Slide": slide_no,
                    "Status": "Updated Combined",
                    "Width": width,
                    "Height": height,
                    "Action": f"{width} X {height}"
                })
            else:
                report.append({
                    "Slide": slide_no,
                    "Status": "Size Not Found",
                    "Width": width,
                    "Height": height,
                    "Action": "No W/X/H size field found"
                })

        out = Path(tempfile.mktemp(suffix=".pptx"))
        prs.save(str(out))

        report_df = pd.DataFrame(report)
        st.dataframe(report_df, use_container_width=True)
        st.download_button(
            "⬇️ Download Fixed PPT",
            out.read_bytes(),
            "fixed_size_ppt.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
        st.download_button(
            "⬇️ Download Processing Report",
            dataframe_download(report_df),
            "size_fixer_report.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

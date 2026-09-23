import io
import openpyxl
import streamlit as st
from pptx import Presentation

def process_delete(excel_file, ppt_file, action_col_name):
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    if action_col_name not in headers:
        raise ValueError(f"Column '{action_col_name}' was not found in the Excel sheet.")
    action_col_idx = headers.index(action_col_name)
    slides_to_delete=[]
    for row_idx,row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        val = str(row[action_col_idx]).strip().lower() if row[action_col_idx] is not None else ""
        if val == "delete": slides_to_delete.append(row_idx-2)
    prs=Presentation(ppt_file)
    valid=[i for i in slides_to_delete if 0 <= i < len(prs.slides)]
    for i in sorted(valid, reverse=True):
        slide_id=prs.slides._sldIdLst[i]
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    out=io.BytesIO(); prs.save(out); out.seek(0)
    return out.getvalue(), len(valid), len(slides_to_delete), len(prs.slides)

def render():
    st.title("🗑️ Excel-Based PPT Slide Deleter")
    st.caption("Excel me 'delete' marked rows ke corresponding PPT slides automatically remove karein.")
    excel_file=st.file_uploader("1. Upload Excel File (.xlsx)", type=["xlsx"], key="deleter_excel")
    ppt_file=st.file_uploader("2. Upload PowerPoint File (.pptx)", type=["pptx"], key="deleter_ppt")
    action_col_name=st.text_input("Excel Action Column Name", value="delete in ppt", key="deleter_action_col")
    if st.button("🚀 Process Presentation", type="primary", use_container_width=True, key="deleter_process"):
        if not excel_file or not ppt_file:
            st.warning("Please upload both Excel and PPT files.")
            return
        try:
            data,deleted,total_marked,remaining=process_delete(excel_file,ppt_file,action_col_name)
            st.success(f"Done — {deleted} slide(s) deleted. {remaining} slide(s) remain.")
            st.download_button("⬇️ Download Updated PPT", data=data, file_name="PPT_Slides_Deleted.pptx", mime="application/vnd.openxmlformats-officedocument.presentationml.presentation", use_container_width=True, key="deleter_download")
        except Exception as e:
            st.error(f"❌ {e}")

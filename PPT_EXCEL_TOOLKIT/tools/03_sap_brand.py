import re
import tempfile
from pathlib import Path
import pandas as pd
import streamlit as st
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from .common import read_excel, save_uploaded, dataframe_download, find_column, normalize


SAP_ALIASES = ["sap code", "sapcode", "sap", "sap no", "sap number", "material code"]
BRAND_ALIASES = ["brand", "brand name"]


def set_text(shape, text):
    if not getattr(shape, "has_text_frame", False):
        return False
    shape.text_frame.clear()
    p = shape.text_frame.paragraphs[0]
    p.text = str(text)
    p.alignment = PP_ALIGN.CENTER
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return True


def find_label_value_shapes(slide, labels):
    for shape in slide.shapes:
        text = getattr(shape, "text", "") or ""
        norm = normalize(text)
        for label in labels:
            if normalize(label) in norm:
                return shape
    return None


def run_tool(max_upload_mb=2048):
    st.subheader("1. Upload files")
    excel = st.file_uploader("Excel Master File", type=["xlsx", "xls", "csv"], key="sap_excel")
    ppt = st.file_uploader("PowerPoint File", type=["pptx"], key="sap_ppt")

    if not excel or not ppt:
        st.info("Upload Excel + PPT to continue.")
        return

    df = read_excel(excel)
    sap_col = find_column(df, SAP_ALIASES)
    brand_col = find_column(df, BRAND_ALIASES)

    if not sap_col and not brand_col:
        st.error("SAP Code ya Brand column nahi mila.")
        st.write("Detected columns:", list(df.columns))
        return

    st.success(f"Excel loaded: {len(df)} rows")
    if st.button("🚀 Transfer SAP + Brand", type="primary", key="sap_start"):
        prs = Presentation(str(save_uploaded(ppt)))
        report = []

        for i, slide in enumerate(prs.slides):
            if i >= len(df):
                break
            row = df.iloc[i]
            sap = "" if not sap_col else str(row[sap_col]) if pd.notna(row[sap_col]) else ""
            brand = "" if not brand_col else str(row[brand_col]) if pd.notna(row[brand_col]) else ""

            changed = []
            for shape in slide.shapes:
                text = getattr(shape, "text", "") or ""
                low = text.lower()
                if sap and ("sap" in low or "sap code" in low):
                    set_text(shape, sap)
                    changed.append("SAP")
                elif brand and "brand" in low:
                    set_text(shape, brand)
                    changed.append("Brand")

            report.append({"Slide": i + 1, "SAP": sap, "Brand": brand, "Updated": ", ".join(changed) or "Not Found"})

        out = Path(tempfile.mktemp(suffix=".pptx"))
        prs.save(str(out))
        report_df = pd.DataFrame(report)
        st.dataframe(report_df, use_container_width=True)
        st.download_button("⬇️ Download PPT", out.read_bytes(), "sap_brand_updated.pptx",
                           "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        st.download_button("⬇️ Download Report", dataframe_download(report_df), "sap_brand_report.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

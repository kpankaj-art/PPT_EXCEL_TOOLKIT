import streamlit as st

import pandas as pd

import re

import io

from pptx import Presentation

from pptx.util import Inches

WIDTH_ALIASES = {
    "w",
    "width",
    "wight",
    "wid",
    "widths",
    "width inch",
    "width inches",
    "width inchs",
    "width in inch",
    "width in inches",
    "width (inch)",
    "width (inches)",
    "w inch",
    "w inches",
    "w (inch)",
    "w (inches)",
    "board width",
    "size width"
}

HEIGHT_ALIASES = {
    "h",
    "height",
    "hight",
    "heigt",
    "ht",
    "height inch",
    "height inches",
    "height inchs",
    "height in inch",
    "height in inches",
    "height (inch)",
    "height (inches)",
    "h inch",
    "h inches",
    "h (inch)",
    "h (inches)",
    "board height",
    "size height"
}

PROTECTED_WORDS = {
    "qty",
    "quantity",
    "media",
    "media type",
    "remarks",
    "remark",
    "outlet",
    "outlet name",
    "address",
    "contact",
    "contact no",
    "contact number",
    "mobile",
    "phone",
    "sap",
    "sap code",
    "outlet code",
    "district",
    "type"
}

def normalize_text(value):

    if value is None:
        return ""

    text = str(value)

    text = text.replace("\n", " ")
    text = text.replace("\r", " ")
    text = text.replace("\t", " ")

    text = text.lower()

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()

def normalize_column_name(value):

    text = normalize_text(value)

    text = text.replace("_", " ")
    text = text.replace("-", " ")
    text = text.replace("/", " ")

    text = text.replace("(", " ")
    text = text.replace(")", " ")

    text = re.sub(
        r"[^a-z0-9\s]",
        " ",
        text
    )

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()

def alias_score(
    column_name,
    aliases,
    kind
):

    normalized = normalize_column_name(
        column_name
    )

    if not normalized:
        return -1

    # -----------------------------------------------------
    # Exact alias
    # -----------------------------------------------------

    if normalized in aliases:

        if normalized in {
            "w",
            "h"
        }:

            return 1000

        if normalized in {
            "width",
            "height"
        }:

            return 1100

        return 1050

    # -----------------------------------------------------
    # Token based
    # -----------------------------------------------------

    words = normalized.split()

    score = 0

    if kind == "width":

        if "width" in words:
            score += 800

        if "wight" in words:
            score += 750

        if "wid" in words:
            score += 700

        if "w" in words:
            score += 600

    elif kind == "height":

        if "height" in words:
            score += 800

        if "hight" in words:
            score += 750

        if "heigt" in words:
            score += 750

        if "ht" in words:
            score += 650

        if "h" in words:
            score += 600

    # -----------------------------------------------------
    # Inch words
    # -----------------------------------------------------

    if "inch" in words:
        score += 30

    if "inches" in words:
        score += 30

    # -----------------------------------------------------
    # Size words
    # -----------------------------------------------------

    if "size" in words:
        score += 10

    return score

def detect_size_columns(
    columns
):

    columns = list(
        columns
    )

    # -----------------------------------------------------
    # Generate scores
    # -----------------------------------------------------

    width_candidates = []
    height_candidates = []

    for column in columns:

        width_score = alias_score(
            column,
            WIDTH_ALIASES,
            "width"
        )

        height_score = alias_score(
            column,
            HEIGHT_ALIASES,
            "height"
        )

        if width_score >= 600:

            width_candidates.append(
                (
                    width_score,
                    column
                )
            )

        if height_score >= 600:

            height_candidates.append(
                (
                    height_score,
                    column
                )
            )

    # Highest score first
    width_candidates.sort(
        key=lambda x: x[0],
        reverse=True
    )

    height_candidates.sort(
        key=lambda x: x[0],
        reverse=True
    )

    # -----------------------------------------------------
    # Find best DIFFERENT pair
    # -----------------------------------------------------

    pair_candidates = []

    for width_score, width_col in (
        width_candidates
    ):

        for height_score, height_col in (
            height_candidates
        ):

            # Same column cannot be Width + Height
            if width_col == height_col:
                continue

            pair_score = (
                width_score
                +
                height_score
            )

            # Strong bonus for exact common pairs
            w_norm = normalize_column_name(
                width_col
            )

            h_norm = normalize_column_name(
                height_col
            )

            # W + H
            if (
                w_norm == "w"
                and
                h_norm == "h"
            ):
                pair_score += 500

            # Width + Height
            if (
                w_norm == "width"
                and
                h_norm == "height"
            ):
                pair_score += 500

            # Wight + Height
            if (
                w_norm == "wight"
                and
                h_norm == "height"
            ):
                pair_score += 450

            pair_candidates.append(
                (
                    pair_score,
                    width_col,
                    height_col
                )
            )

    if pair_candidates:

        pair_candidates.sort(
            key=lambda x: x[0],
            reverse=True
        )

        best = pair_candidates[0]

        return (
            best[1],
            best[2]
        )

    return (
        None,
        None
    )

def clean_number(value):

    if pd.isna(value):
        return ""

    text = str(value).strip()

    if not text:
        return ""

    text = text.replace(
        ",",
        ""
    )

    # Remove units
    text = re.sub(
        r"\s*(inch|inches|in)\s*$",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = text.strip()

    try:

        number = float(text)

        if number.is_integer():

            return str(
                int(number)
            )

        return (
            str(number)
            .rstrip("0")
            .rstrip(".")
        )

    except Exception:

        return text

def get_shape_text(
    shape
):

    try:

        if not shape.has_text_frame:
            return ""

        return shape.text or ""

    except Exception:

        return ""

def set_shape_text(
    shape,
    new_text
):

    try:

        if not shape.has_text_frame:
            return False

        text_frame = shape.text_frame

        # Preserve first run formatting
        if text_frame.paragraphs:

            first_paragraph = (
                text_frame.paragraphs[0]
            )

            if first_paragraph.runs:

                first_run = (
                    first_paragraph.runs[0]
                )

                first_run.text = str(
                    new_text
                )

                # Clear extra runs
                for run in (
                    first_paragraph.runs[1:]
                ):

                    run.text = ""

                # Clear extra paragraphs
                for paragraph in (
                    text_frame.paragraphs[1:]
                ):

                    for run in paragraph.runs:

                        run.text = ""

                return True

        text_frame.text = str(
            new_text
        )

        return True

    except Exception:

        try:

            shape.text = str(
                new_text
            )

            return True

        except Exception:

            return False

def collect_text_shapes(
    slide
):

    items = []

    for index, shape in enumerate(
        slide.shapes
    ):

        text = get_shape_text(
            shape
        )

        if not text.strip():
            continue

        items.append({
            "index": index,
            "shape": shape,
            "text": text,
            "norm": normalize_text(text)
        })

    return items

def contains_size_word(
    text
):

    if not text:
        return False

    return bool(
        re.search(
            r"\bsize\b",
            str(text),
            flags=re.IGNORECASE
        )
    )

def contains_full_size(
    text
):

    if not text:
        return False

    return bool(
        FULL_SIZE_PATTERN.search(
            str(text)
        )
    )

def replace_size_inside_text(
    text,
    width,
    height
):

    if not text:
        return text, False

    original = str(
        text
    )

    # Find Size
    size_match = re.search(
        r"\bsize\b",
        original,
        flags=re.IGNORECASE
    )

    if not size_match:

        return original, False

    after_size_start = (
        size_match.end()
    )

    after_size = original[
        after_size_start:
    ]

    # Find dimension AFTER Size
    size_value_match = (
        FULL_SIZE_PATTERN.search(
            after_size
        )
    )

    if not size_value_match:

        return original, False

    start = (
        after_size_start
        +
        size_value_match.start()
    )

    end = (
        after_size_start
        +
        size_value_match.end()
    )

    new_value = (
        f"{width}X{height}"
    )

    new_text = (
        original[:start]
        +
        new_value
        +
        original[end:]
    )

    return new_text, True

def find_inline_size_shapes(
    items
):

    result = []

    for item in items:

        text = item[
            "text"
        ]

        if not contains_size_word(
            text
        ):
            continue

        if not contains_full_size(
            text
        ):
            continue

        result.append(
            item
        )

    return result

def find_size_labels(
    items
):

    result = []

    for item in items:

        text = item[
            "text"
        ]

        if not contains_size_word(
            text
        ):
            continue

        # Inline size already handled
        if contains_full_size(
            text
        ):
            continue

        result.append(
            item
        )

    return result

def is_number_text(
    text
):

    if not text:
        return False

    return bool(
        re.fullmatch(
            r"\s*\d+(?:\.\d+)?\s*",
            str(text)
        )
    )

def is_x_text(
    text
):

    if not text:
        return False

    return normalize_text(
        text
    ) in {
        "x",
        "×",
        "*"
    }

def find_separate_size(
    items,
    labels
):

    candidates = []

    for label in labels:

        label_shape = label[
            "shape"
        ]

        label_left = (
            label_shape.left
        )

        label_right = (
            label_shape.left
            +
            label_shape.width
        )

        label_center_y = (
            label_shape.top
            +
            label_shape.height / 2
        )

        # -------------------------------------------------
        # Find X
        # -------------------------------------------------

        x_candidates = []

        for item in items:

            shape = item[
                "shape"
            ]

            if shape is label_shape:
                continue

            if not is_x_text(
                item["text"]
            ):
                continue

            center_x = (
                shape.left
                +
                shape.width / 2
            )

            center_y = (
                shape.top
                +
                shape.height / 2
            )

            if (
                center_x <
                label_left - Inches(0.5)
            ):
                continue

            if (
                center_x >
                label_right + Inches(0.5)
            ):
                continue

            if abs(
                center_y -
                label_center_y
            ) > Inches(0.7):

                continue

            x_candidates.append(
                item
            )

        # -------------------------------------------------
        # Find Width + Height around X
        # -------------------------------------------------

        for x_item in x_candidates:

            x_shape = x_item[
                "shape"
            ]

            x_center_x = (
                x_shape.left
                +
                x_shape.width / 2
            )

            x_center_y = (
                x_shape.top
                +
                x_shape.height / 2
            )

            left_numbers = []
            right_numbers = []

            for item in items:

                shape = item[
                    "shape"
                ]

                if shape is label_shape:
                    continue

                if shape is x_shape:
                    continue

                if not is_number_text(
                    item["text"]
                ):
                    continue

                center_x = (
                    shape.left
                    +
                    shape.width / 2
                )

                center_y = (
                    shape.top
                    +
                    shape.height / 2
                )

                if abs(
                    center_y -
                    x_center_y
                ) > Inches(0.55):

                    continue

                # Stay around Size field
                if (
                    center_x <
                    label_left - Inches(0.25)
                ):
                    continue

                if (
                    center_x >
                    label_right + Inches(0.25)
                ):
                    continue

                if center_x < x_center_x:

                    left_numbers.append(
                        item
                    )

                elif center_x > x_center_x:

                    right_numbers.append(
                        item
                    )

            if (
                not left_numbers
                or
                not right_numbers
            ):
                continue

            left_numbers.sort(
                key=lambda item:
                abs(
                    (
                        item["shape"].left
                        +
                        item["shape"].width / 2
                    )
                    -
                    x_center_x
                )
            )

            right_numbers.sort(
                key=lambda item:
                abs(
                    (
                        item["shape"].left
                        +
                        item["shape"].width / 2
                    )
                    -
                    x_center_x
                )
            )

            candidates.append({
                "label": label,
                "width": left_numbers[0],
                "x": x_item,
                "height": right_numbers[0]
            })

    if not candidates:
        return None

    return candidates[0]

def find_standalone_size(
    items,
    labels
):

    for item in items:

        text = item[
            "text"
        ].strip()

        if not re.fullmatch(
            r"\d+(?:\.\d+)?\s*[x×*]\s*\d+(?:\.\d+)?",
            text,
            flags=re.IGNORECASE
        ):
            continue

        shape = item[
            "shape"
        ]

        center_x = (
            shape.left
            +
            shape.width / 2
        )

        center_y = (
            shape.top
            +
            shape.height / 2
        )

        for label in labels:

            label_shape = label[
                "shape"
            ]

            label_center_x = (
                label_shape.left
                +
                label_shape.width / 2
            )

            label_center_y = (
                label_shape.top
                +
                label_shape.height / 2
            )

            if abs(
                center_y -
                label_center_y
            ) > Inches(0.7):

                continue

            if abs(
                center_x -
                label_center_x
            ) > Inches(4):

                continue

            return item

    return None

def process_slide(
    slide,
    width,
    height
):

    items = collect_text_shapes(
        slide
    )

    # =====================================================
    # 1. SIZE INSIDE SAME TEXTBOX
    # =====================================================

    inline_items = find_inline_size_shapes(
        items
    )

    for item in inline_items:

        old_text = item[
            "text"
        ]

        new_text, changed = (
            replace_size_inside_text(
                old_text,
                width,
                height
            )
        )

        if changed:

            success = set_shape_text(
                item["shape"],
                new_text
            )

            if success:

                return {
                    "status": "Updated Size Inside Textbox",
                    "action": (
                        f"{old_text} -> {new_text}"
                    )
                }

    # =====================================================
    # 2. SEPARATE SIZE FIELDS
    # =====================================================

    labels = find_size_labels(
        items
    )

    separate = find_separate_size(
        items,
        labels
    )

    if separate is not None:

        old_width = separate[
            "width"
        ]["text"]

        old_height = separate[
            "height"
        ]["text"]

        ok_width = set_shape_text(
            separate["width"]["shape"],
            str(width)
        )

        ok_x = set_shape_text(
            separate["x"]["shape"],
            "X"
        )

        ok_height = set_shape_text(
            separate["height"]["shape"],
            str(height)
        )

        if (
            ok_width
            and
            ok_x
            and
            ok_height
        ):

            return {
                "status": "Updated Separate Size Fields",
                "action": (
                    f"{old_width} X {old_height}"
                    f" -> "
                    f"{width} X {height}"
                )
            }

    # =====================================================
    # 3. EXISTING COMBINED SIZE BOX
    # =====================================================

    standalone = find_standalone_size(
        items,
        labels
    )

    if standalone is not None:

        old_text = standalone[
            "text"
        ]

        new_text = (
            f"{width} X {height}"
        )

        success = set_shape_text(
            standalone["shape"],
            new_text
        )

        if success:

            return {
                "status": "Updated Existing Size Box",
                "action": (
                    f"{old_text} -> {new_text}"
                )
            }

    # =====================================================
    # 4. NOT FOUND
    # =====================================================

    return {
        "status": "Size Not Found",
        "action": (
            "Size field/value detect nahi hua. "
            "New box create nahi kiya."
        )
    }

def read_excel_file(
    uploaded_file
):

    filename = (
        uploaded_file.name.lower()
    )

    if filename.endswith(
        ".csv"
    ):

        return {
            "CSV": pd.read_csv(
                uploaded_file
            )
        }

    if filename.endswith(
        ".xls"
    ):

        return pd.read_excel(
            uploaded_file,
            sheet_name=None,
            engine="xlrd"
        )

    return pd.read_excel(
        uploaded_file,
        sheet_name=None,
        engine="openpyxl"
    )

def get_dataframe(
    uploaded_file
):

    sheets = read_excel_file(
        uploaded_file
    )

    if isinstance(
        sheets,
        pd.DataFrame
    ):

        return sheets, "CSV"

    # Merged_Result priority
    if "Merged_Result" in sheets:

        return (
            sheets["Merged_Result"],
            "Merged_Result"
        )

    # Case insensitive
    for name, df in sheets.items():

        if normalize_text(
            name
        ) == "merged_result":

            return df, name

    # First non-empty sheet
    for name, df in sheets.items():

        if (
            isinstance(
                df,
                pd.DataFrame
            )
            and
            not df.empty
        ):

            return df, name

    first_name = list(
        sheets.keys()
    )[0]

    return (
        sheets[first_name],
        first_name
    )

def process_ppt(
    ppt_bytes,
    df,
    width_col,
    height_col
):

    prs = Presentation(
        io.BytesIO(
            ppt_bytes
        )
    )

    total_slides = len(
        prs.slides
    )

    results = []

    # -----------------------------------------------------
    # Prepare Excel rows
    # -----------------------------------------------------

    excel_rows = []

    for index, row in df.iterrows():

        width = clean_number(
            row.get(
                width_col,
                ""
            )
        )

        height = clean_number(
            row.get(
                height_col,
                ""
            )
        )

        excel_rows.append({
            "excel_row": index + 2,
            "width": width,
            "height": height
        })

    progress = st.progress(
        0
    )

    status_box = st.empty()

    # =====================================================
    # SLIDES
    # =====================================================

    for slide_number, slide in enumerate(
        prs.slides,
        start=1
    ):

        status_box.write(
            f"Processing Slide {slide_number} / {total_slides}"
        )

        excel_index = (
            slide_number - 1
        )

        # -------------------------------------------------
        # Excel row unavailable
        # -------------------------------------------------

        if excel_index >= len(
            excel_rows
        ):

            results.append({
                "Slide": slide_number,
                "Excel Row": "",
                "Width": "",
                "Height": "",
                "Status": "Excel Row Not Available",
                "Action": ""
            })

            progress.progress(
                slide_number /
                total_slides
            )

            continue

        row = excel_rows[
            excel_index
        ]

        width = row[
            "width"
        ]

        height = row[
            "height"
        ]

        # -------------------------------------------------
        # Missing size
        # -------------------------------------------------

        if (
            width == ""
            or
            height == ""
        ):

            results.append({
                "Slide": slide_number,
                "Excel Row": row[
                    "excel_row"
                ],
                "Width": width,
                "Height": height,
                "Status": "Width/Height Missing",
                "Action": ""
            })

            progress.progress(
                slide_number /
                total_slides
            )

            continue

        # -------------------------------------------------
        # Process slide
        # -------------------------------------------------

        result = process_slide(
            slide,
            width,
            height
        )

        results.append({
            "Slide": slide_number,
            "Excel Row": row[
                "excel_row"
            ],
            "Width": width,
            "Height": height,
            "Status": result[
                "status"
            ],
            "Action": result[
                "action"
            ]
        })

        progress.progress(
            slide_number /
            total_slides
        )

    progress.progress(
        1.0
    )

    status_box.success(
        f"Completed {total_slides} slides."
    )

    return (
        prs,
        pd.DataFrame(
            results
        )
    )


def render():
    st.title("📐 PPT SIZE FIXER")

    st.caption(
        "Excel ke Width/Height columns automatically detect karke "
        "PPT ke existing Size field ko update karega."
    )

    if "fixed_ppt_bytes" not in st.session_state:
        st.session_state.fixed_ppt_bytes = None

    if "result_df" not in st.session_state:
        st.session_state.result_df = None

    if "output_filename" not in st.session_state:
        st.session_state.output_filename = (
            "PPT_SIZE_FIXED_V9.pptx"
        )

    FULL_SIZE_PATTERN = re.compile(
        r"""
        (?P<width>
            \d+(?:\.\d+)?
        )
        \s*
        [x×*]
        \s*
        (?P<height>
            \d+(?:\.\d+)?
        )
        """,
        re.IGNORECASE |
        re.VERBOSE
    )

    st.markdown(
        "### 📊 Excel Master File"
    )

    excel_file = st.file_uploader(
        "Upload Excel",
        type=[
            "xlsx",
            "xls",
            "xlsm",
            "csv"
        ],
        key="excel_upload_v9"
    )

    st.markdown(
        "### 📄 PowerPoint File"
    )

    ppt_file = st.file_uploader(
        "Upload PPTX",
        type=[
            "pptx"
        ],
        key="ppt_upload_v9"
    )

    if excel_file is not None:

        try:

            df, sheet_name = get_dataframe(
                excel_file
            )

            # =================================================
            # AUTOMATIC WIDTH + HEIGHT DETECTION
            # =================================================

            width_col, height_col = (
                detect_size_columns(
                    df.columns
                )
            )

            # -------------------------------------------------
            # Width/Height not detected
            # -------------------------------------------------

            if (
                width_col is None
                or
                height_col is None
            ):

                st.error(
                    "❌ Width/Height columns automatically detect nahi hue."
                )

                st.write(
                    "Excel ke available headings:"
                )

                st.write(
                    list(df.columns)
                )

                st.stop()

            # -------------------------------------------------
            # Same column safety
            # -------------------------------------------------

            if width_col == height_col:

                st.error(
                    "❌ Width aur Height same column detect hue. "
                    "Processing stop kar di gayi."
                )

                st.write(
                    "Available headings:"
                )

                st.write(
                    list(df.columns)
                )

                st.stop()

            # -------------------------------------------------
            # Small confirmation
            # -------------------------------------------------

            st.success(
                f"Excel loaded successfully | "
                f"Width = {width_col} | "
                f"Height = {height_col}"
            )

            # =================================================
            # PPT
            # =================================================

            if ppt_file is not None:

                if st.button(
                    "🚀 FIX PPT SIZE",
                    type="primary",
                    use_container_width=True
                ):

                    try:

                        # Clear old result
                        st.session_state.fixed_ppt_bytes = None
                        st.session_state.result_df = None

                        ppt_bytes = (
                            ppt_file.getvalue()
                        )

                        # -------------------------------------
                        # Process
                        # -------------------------------------

                        prs, result_df = process_ppt(
                            ppt_bytes,
                            df,
                            width_col,
                            height_col
                        )

                        # -------------------------------------
                        # Save
                        # -------------------------------------

                        output_buffer = io.BytesIO()

                        prs.save(
                            output_buffer
                        )

                        output_buffer.seek(
                            0
                        )

                        st.session_state.fixed_ppt_bytes = (
                            output_buffer.getvalue()
                        )

                        st.session_state.result_df = (
                            result_df
                        )

                        st.session_state.output_filename = (
                            "PPT_SIZE_FIXED_V9.pptx"
                        )

                        st.success(
                            "✅ PPT successfully processed."
                        )

                    except Exception as e:

                        st.error(
                            f"❌ Processing Error: {str(e)}"
                        )

                        st.exception(e)

            else:

                st.info(
                    "PPTX file upload karo."
                )

        except Exception as e:

            st.error(
                f"❌ Excel Error: {str(e)}"
            )

            st.exception(e)

    else:

        st.info(
            "Pehle Excel Master File upload karo."
        )

    if (
        st.session_state.fixed_ppt_bytes
        is not None
    ):

        st.markdown(
            "---"
        )

        st.markdown(
            "### ✅ Fixed PPT Ready"
        )

        st.download_button(
            label="⬇️ DOWNLOAD FIXED PPT",
            data=st.session_state.fixed_ppt_bytes,
            file_name=st.session_state.output_filename,
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation"
            ),
            type="primary",
            use_container_width=True,
            on_click="ignore"
        )

    if (
        st.session_state.result_df
        is not None
    ):

        result_df = (
            st.session_state.result_df
        )

        st.markdown(
            "### 📋 Processing Report"
        )

        updated = len(
            result_df[
                result_df[
                    "Status"
                ].str.contains(
                    "Updated",
                    na=False
                )
            ]
        )

        not_found = len(
            result_df[
                result_df[
                    "Status"
                ].str.contains(
                    "Not Found",
                    na=False
                )
            ]
        )

        missing = len(
            result_df[
                result_df[
                    "Status"
                ].str.contains(
                    "Missing",
                    na=False
                )
            ]
        )

        c1, c2, c3 = st.columns(
            3
        )

        with c1:

            st.metric(
                "Updated",
                updated
            )

        with c2:

            st.metric(
                "Size Not Found",
                not_found
            )

        with c3:

            st.metric(
                "Width/Height Missing",
                missing
            )

        st.dataframe(
            result_df,
            use_container_width=True,
            height=500
        )

    st.markdown(
        "---"
    )

    st.caption(
        "PPT SIZE FIXER V9 | "
        "Excel Row 2 → Slide 1 | "
        "Excel Row 3 → Slide 2 | "
        "Automatic W/H detection | "
        "Existing Size field only"
    )

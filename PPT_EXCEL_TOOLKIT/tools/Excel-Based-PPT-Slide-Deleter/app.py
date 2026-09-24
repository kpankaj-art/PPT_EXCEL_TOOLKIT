import streamlit as st
import openpyxl
from pptx import Presentation
import io

# Page Configuration
st.set_page_config(page_title="Excel to PPT Cleaner", page_icon="📊", layout="centered")

st.title("📊 Excel-Based PPT Slide Deleter")
st.write(
    "Upload your Excel and PowerPoint files. Slides corresponding to rows marked with **delete** "
    "under the specified column will be automatically removed."
)

# File Uploaders
excel_file = st.file_uploader("1. Upload Excel File (.xlsx)", type=["xlsx"])
ppt_file = st.file_uploader("2. Upload PowerPoint File (.pptx)", type=["pptx"])

# Column Name Configuration
action_col_name = st.text_input("Excel Action Column Name", value="delete in ppt")

if excel_file and ppt_file:
    if st.button("Process Presentation 🚀", type="primary"):
        try:
            # 1. Read Excel File
            wb = openpyxl.load_workbook(excel_file)
            sheet = wb.active

            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
            
            if action_col_name not in headers:
                st.error(f"❌ Error: Column '{action_col_name}' was not found in the Excel sheet.")
            else:
                action_col_idx = headers.index(action_col_name) + 1

                # Find slides marked for deletion (0-indexed base for PPT)
                slides_to_delete = []
                for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                    action_val = str(row[action_col_idx - 1]).strip().lower() if row[action_col_idx - 1] is not None else ""
                    if action_val == "delete":
                        # Row 2 -> Slide 1 (Index 0)
                        slides_to_delete.append(row_idx - 2)

                st.info(f"Total slides marked for deletion: {len(slides_to_delete)}")

                # 2. Process PPT File
                prs = Presentation(ppt_file)
                total_slides = len(prs.slides)

                # Delete in reverse order to prevent index shifting
                for slide_idx in sorted(slides_to_delete, reverse=True):
                    if slide_idx < total_slides:
                        rId = prs.slides._sldIdLst[slide_idx].rId
                        prs.part.drop_rel(rId)
                        del prs.slides._sldIdLst[slide_idx]

                # 3. Save PPT to Memory Stream for Direct Download
                output_stream = io.BytesIO()
                prs.save(output_stream)
                output_stream.seek(0)

                st.success("✅ Presentation updated successfully!")
                
                # Download Button
                st.download_button(
                    label="📥 Download Updated PPT",
                    data=output_stream,
                    file_name="Updated_Presentation.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                )

        except Exception as e:
            st.error(f"Execution Error: {str(e)}")

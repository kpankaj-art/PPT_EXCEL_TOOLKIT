import tempfile
from pathlib import Path
import pandas as pd
import streamlit as st
from pptx import Presentation
from .common import read_excel, save_uploaded, find_column, dataframe_download


SLIDE_ALIASES = ["slide", "slide no", "slide number", "slide id", "page", "page no"]


def run_tool(max_upload_mb=2048):
    st.subheader("1. Upload files")
    excel = st.file_uploader("Excel File with Slide Numbers", type=["xlsx", "xls", "csv"], key="del_excel")
    ppt = st.file_uploader("PowerPoint File", type=["pptx"], key="del_ppt")
    if not excel or not ppt:
        st.info("Upload Excel + PPT to continue.")
        return

    df = read_excel(excel)
    slide_col = find_column(df, SLIDE_ALIASES)
    if not slide_col:
        st.error("Slide number column nahi mila.")
        st.write("Detected columns:", list(df.columns))
        return

    if st.button("🗑️ Delete Slides", type="primary", key="del_start"):
        numbers = set()
        for value in df[slide_col].dropna():
            try:
                numbers.add(int(float(value)))
            except Exception:
                pass

        prs = Presentation(str(save_uploaded(ppt)))
        existing = len(prs.slides)
        deleted = []
        for slide_no in sorted(numbers, reverse=True):
            if 1 <= slide_no <= len(prs.slides):
                slide = prs.slides[slide_no - 1]
                rel_id = slide.part.rels.get(r"http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout")
                slide_id = prs.slides._sldIdLst[slide_no - 1]
                prs.slides._sldIdLst.remove(slide_id)
                deleted.append(slide_no)

        out = Path(tempfile.mktemp(suffix=".pptx"))
        prs.save(str(out))
        report = pd.DataFrame([{
            "Original Slides": existing,
            "Requested Slides": len(numbers),
            "Deleted Slides": len(deleted),
            "Remaining Slides": len(prs.slides),
            "Deleted Slide Numbers": ", ".join(map(str, sorted(deleted))) or "None"
        }])
        st.dataframe(report, use_container_width=True)
        st.download_button("⬇️ Download PPT", out.read_bytes(), "slides_deleted.pptx",
                           "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        st.download_button("⬇️ Download Report", dataframe_download(report), "slide_delete_report.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
